import hashlib
import json
import re
from collections.abc import Sequence
from dataclasses import dataclass
from datetime import UTC, datetime
from decimal import Decimal
from pathlib import Path
from typing import Protocol
from uuid import UUID, uuid4

from openpyxl import Workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from app.modules.selection_export.domain import FrozenSnapshot, GeneratedExportFile

ExcelValue = str | float | datetime | None


@dataclass(frozen=True)
class ExportColumn:
    field_key: str
    display_name: str
    width: float
    value_type: str = "text"


class StageExporter(Protocol):
    schema_version: str

    def export(
        self,
        snapshot: FrozenSnapshot,
        *,
        requested_by: UUID,
        requested_at: datetime,
        store_names: Sequence[str],
        output_directory: Path,
    ) -> GeneratedExportFile: ...


class InitialSelectionExporter:
    """Export only stable fields until the real Skill output schema is registered."""

    schema_version = "initial_selection.v1"
    sheet_name = "选品表"
    file_prefix = "AI选品表"
    columns: tuple[ExportColumn, ...] = (
        ExportColumn("selection_id", "选品记录ID", 38),
        ExportColumn("store_id", "店铺ID", 24),
        ExportColumn("marketplace", "平台", 18),
        ExportColumn("product_name", "商品名称", 36),
        ExportColumn("source_type", "来源类型", 18),
        ExportColumn("source_ref", "来源引用", 30),
        ExportColumn("source_snapshot_ref", "来源快照引用", 30),
        ExportColumn("source_version", "来源版本", 18),
        ExportColumn("owner_id", "负责人ID", 38),
        ExportColumn("stage", "当前阶段", 18),
        ExportColumn("revision", "记录修订号", 14, "integer"),
        ExportColumn("created_at", "创建时间", 21, "datetime"),
        ExportColumn("updated_at", "更新时间", 21, "datetime"),
    )

    def export(
        self,
        snapshot: FrozenSnapshot,
        *,
        requested_by: UUID,
        requested_at: datetime,
        store_names: Sequence[str],
        output_directory: Path,
    ) -> GeneratedExportFile:
        output_directory.mkdir(parents=True, exist_ok=True)
        file_name = self._file_name(snapshot, requested_at, store_names)
        target = (output_directory / f"{snapshot.snapshot_id}_{file_name}").resolve()
        if target.parent != output_directory.resolve():
            raise ValueError("Unsafe export file path")
        temporary = target.with_name(f".{target.name}.{uuid4().hex}.tmp")
        workbook = Workbook(write_only=True)
        self._write_data_sheet(workbook, snapshot)
        self._write_information_sheet(workbook, snapshot, requested_by, requested_at)
        try:
            workbook.save(temporary)
            temporary.replace(target)
        finally:
            temporary.unlink(missing_ok=True)
        checksum = hashlib.sha256(target.read_bytes()).hexdigest()
        return GeneratedExportFile(path=target, file_name=file_name, checksum=checksum)

    def business_rows(self, snapshot: FrozenSnapshot) -> list[list[ExcelValue]]:
        return [
            [
                self._typed_value(row.rendered_row.get(column.field_key), column)
                for column in self.columns
            ]
            for row in snapshot.rows
        ]

    def _write_data_sheet(self, workbook: Workbook, snapshot: FrozenSnapshot) -> None:
        sheet = workbook.create_sheet(self.sheet_name)
        sheet.freeze_panes = "A2"
        for index, column in enumerate(self.columns, start=1):
            sheet.column_dimensions[get_column_letter(index)].width = column.width
        header = []
        for column in self.columns:
            cell = WriteOnlyCell(sheet, value=column.display_name)
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="1F4E78")
            header.append(cell)
        sheet.append(header)
        for values in self.business_rows(snapshot):
            cells = []
            for value, column in zip(values, self.columns, strict=True):
                cell = WriteOnlyCell(sheet, value=value)
                if column.value_type == "datetime" and value is not None:
                    cell.number_format = "yyyy-mm-dd hh:mm:ss"
                elif column.value_type == "decimal" and value is not None:
                    cell.number_format = "0.0000"
                cells.append(cell)
            sheet.append(cells)
        sheet.auto_filter.ref = f"A1:{get_column_letter(len(self.columns))}{snapshot.row_count + 1}"

    def _write_information_sheet(
        self,
        workbook: Workbook,
        snapshot: FrozenSnapshot,
        requested_by: UUID,
        requested_at: datetime,
    ) -> None:
        sheet = workbook.create_sheet("导出说明")
        sheet.column_dimensions["A"].width = 28
        sheet.column_dimensions["B"].width = 90
        items: tuple[tuple[str, ExcelValue], ...] = (
            ("导出阶段", snapshot.stage.value),
            ("快照编号", str(snapshot.snapshot_id)),
            ("结构版本", snapshot.schema_version),
            ("导出人", str(requested_by)),
            ("导出时间", requested_at.astimezone(UTC).replace(tzinfo=None)),
            ("店铺范围", ", ".join(snapshot.store_scope) or "空范围"),
            (
                "筛选条件",
                json.dumps(snapshot.filters, ensure_ascii=False, sort_keys=True),
            ),
            ("行数", snapshot.row_count),
            (
                "数据状态说明",
                (
                    "来自不可变整表快照; 空白不是数字0。"
                    "动态Skill表头尚未登记, 本版只导出已确认系统字段。"
                ),
            ),
            (
                "文件校验和",
                f"sha256(snapshot-content):{snapshot.snapshot_checksum}",
            ),
            *self._additional_information(snapshot),
        )
        for label, value in items:
            label_cell = WriteOnlyCell(sheet, value=label)
            label_cell.font = Font(bold=True)
            value_cell = WriteOnlyCell(sheet, value=value)
            if label == "导出时间":
                value_cell.number_format = "yyyy-mm-dd hh:mm:ss"
            sheet.append([label_cell, value_cell])

    @staticmethod
    def _typed_value(value: object, column: ExportColumn) -> ExcelValue:
        if value is None:
            return None
        if column.value_type == "integer":
            if isinstance(value, bool):
                raise ValueError(f"Invalid integer value for {column.field_key}")
            if isinstance(value, (int, str)):
                return int(value)
            raise ValueError(f"Invalid integer value for {column.field_key}")
        if column.value_type == "decimal":
            if isinstance(value, bool):
                raise ValueError(f"Invalid decimal value for {column.field_key}")
            return float(Decimal(str(value)))
        if column.value_type == "datetime":
            parsed = value if isinstance(value, datetime) else datetime.fromisoformat(str(value))
            if parsed.tzinfo is not None:
                parsed = parsed.astimezone(UTC).replace(tzinfo=None)
            return parsed
        if isinstance(value, str) and value.startswith(("=", "+", "-", "@")):
            return f"'{value}"
        return str(value) if not isinstance(value, (int, float)) else value

    @classmethod
    def _file_name(
        cls,
        snapshot: FrozenSnapshot,
        requested_at: datetime,
        store_names: Sequence[str],
    ) -> str:
        scope_label = store_names[0] if len(store_names) == 1 else "多店铺"
        if not store_names:
            scope_label = "空范围"
        safe_scope = re.sub(r"[<>:\"/\\|?*\x00-\x1f]", "_", scope_label).strip(" ._")
        safe_scope = safe_scope[:60] or "未命名店铺"
        date = requested_at.astimezone(UTC).strftime("%Y%m%d")
        return f"{cls.file_prefix}_{safe_scope}_{date}_{str(snapshot.snapshot_id)[:8]}.xlsx"

    @staticmethod
    def _additional_information(
        snapshot: FrozenSnapshot,
    ) -> tuple[tuple[str, ExcelValue], ...]:
        return ()


class QuotationExporter(InitialSelectionExporter):
    schema_version = "quotation.v1"
    sheet_name = "采购询价表"
    file_prefix = "采购询价表"
    columns = (
        ExportColumn("selection_id", "选品记录ID", 38),
        ExportColumn("product_name", "商品名称", 36),
        ExportColumn("store_id", "店铺ID", 24),
        ExportColumn("marketplace", "平台/站点", 18),
        ExportColumn("supplier_name", "供应商", 30),
        ExportColumn("unit_cost", "报价", 15, "decimal"),
        ExportColumn("currency", "币种", 10),
        ExportColumn("moq", "最小起订量", 14, "integer"),
        ExportColumn("freight_per_unit", "单件运费", 15, "decimal"),
        ExportColumn("lead_time_days", "交期(天)", 12, "integer"),
        ExportColumn("value_status", "数据状态", 15),
        ExportColumn("quotation_version", "询价版本", 12, "integer"),
        ExportColumn("submission_status", "提交状态", 12),
        ExportColumn("recorded_by", "登记人ID", 38),
        ExportColumn("quotation_created_at", "登记时间", 21, "datetime"),
        ExportColumn("source_version", "选品来源版本", 18),
    )

    @staticmethod
    def _additional_information(
        snapshot: FrozenSnapshot,
    ) -> tuple[tuple[str, ExcelValue], ...]:
        source_versions = sorted(
            {
                str(row.rendered_row["source_version"])
                for row in snapshot.rows
                if row.rendered_row.get("source_version") is not None
            }
        )
        quotation_versions = sorted(
            {
                int(str(row.rendered_row["quotation_version"]))
                for row in snapshot.rows
                if row.rendered_row.get("quotation_version") is not None
            }
        )
        currencies = sorted(
            {
                str(row.rendered_row["currency"])
                for row in snapshot.rows
                if row.rendered_row.get("currency") is not None
            }
        )
        version_range = (
            f"{quotation_versions[0]}-{quotation_versions[-1]}"
            if quotation_versions
            else "无询价行"
        )
        return (
            ("选品来源版本", ", ".join(source_versions) or "无询价行"),
            ("询价版本范围", version_range),
            ("草稿/正式状态", "正式: 仅包含 submission_status=submitted"),
            (
                "币种说明",
                (
                    f"币种: {', '.join(currencies)}; 报价和单件运费按各行币种解释"
                    if currencies
                    else "无询价行"
                ),
            ),
        )
