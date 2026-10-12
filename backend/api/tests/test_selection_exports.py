from collections.abc import AsyncIterator, Sequence
from datetime import UTC, datetime
from io import BytesIO
from pathlib import Path
from uuid import UUID, uuid4

import httpx
import pytest
import pytest_asyncio
from fastapi import FastAPI
from openpyxl import load_workbook
from pydantic import SecretStr
from sqlalchemy import func, select, update

from app.core.config import Settings
from app.db.base import Base
from app.db.models import (
    SelectionQuotationVersion,
    SelectionRecord,
    SelectionTableSnapshot,
    SelectionTableSnapshotRow,
    StoreRecord,
)
from app.main import create_app
from app.modules.admin.repository import InMemoryAdminRepository
from app.modules.audit.repository import InMemoryAuditRepository
from app.modules.identity.domain import ROLE_PERMISSIONS, Permission, RoleName, UserIdentity
from app.modules.identity.repository import InMemoryIdentityRepository
from app.modules.selection_export.domain import (
    FrozenSnapshot,
    FrozenSnapshotRow,
    GeneratedExportFile,
    SelectionExportStage,
)
from app.modules.selection_export.exporter import (
    ExportColumn,
    InitialSelectionExporter,
    QuotationExporter,
)
from app.modules.store.repository import InMemoryStoreRepository
from app.modules.workspace.repository import InMemoryWorkspaceRepository
from tests.conftest import PASSWORD, USER_IDS, build_users


class FailingInitialSelectionExporter:
    schema_version = "initial_selection.v1"

    def export(
        self,
        snapshot: FrozenSnapshot,
        *,
        requested_by: UUID,
        requested_at: datetime,
        store_names: Sequence[str],
        output_directory: Path,
    ) -> GeneratedExportFile:
        raise RuntimeError("secret=must-not-escape")


def _export_users() -> list[UserIdentity]:
    users: list[UserIdentity] = []
    for user in build_users():
        role = next(iter(user.roles))
        if role in {RoleName.PROCUREMENT, RoleName.OPERATIONS}:
            direct = frozenset(
                {Permission.SELECTION_EXPORT, Permission.QUOTATION_EXPORT}
            )
            user = user.model_copy(
                update={
                    "direct_permissions": direct,
                    "permissions": user.permissions.union(direct),
                }
            )
        users.append(user)
    return users


@pytest_asyncio.fixture
async def export_app(tmp_path: Path) -> AsyncIterator[tuple[FastAPI, InMemoryAuditRepository]]:
    audit_repository = InMemoryAuditRepository()
    settings = Settings(
        app_env="test",
        database_url=f"sqlite+aiosqlite:///{tmp_path / 'exports.db'}",
        auth_secret_key=SecretStr("test-secret-with-sufficient-random-looking-length"),
        export_storage_path=tmp_path / "private-exports",
    )
    app = create_app(
        settings,
        identity_repository=InMemoryIdentityRepository(_export_users()),
        audit_repository=audit_repository,
        admin_repository=InMemoryAdminRepository(),
        workspace_repository=InMemoryWorkspaceRepository(),
        store_repository=InMemoryStoreRepository(),
    )
    async with app.state.engine.begin() as connection:
        await connection.run_sync(Base.metadata.create_all)
    async with app.state.session_factory.begin() as session:
        session.add_all(
            [
                StoreRecord(
                    store_id="store_procurement",
                    display_name="采购演示店铺",
                    marketplace="demo",
                    organization_id="org_1",
                    status="active",
                    revision=1,
                ),
                StoreRecord(
                    store_id="store_other",
                    display_name="其他店铺",
                    marketplace="demo",
                    organization_id="org_2",
                    status="active",
                    revision=1,
                ),
            ]
        )
        regular = SelectionRecord(
                    store_id="store_procurement",
                    marketplace="amazon-us",
                    product_name="Regular product",
                    source_type="skill",
                    source_ref=None,
                    source_snapshot_ref="snapshot://source-1",
                    source_version="skill-v1",
                    stage="sourced",
                    revision=1,
                )
        dangerous = SelectionRecord(
                    store_id="store_procurement",
                    marketplace="amazon-us",
                    product_name="=DANGEROUS()",
                    source_type="skill",
                    source_ref="task://source-2",
                    source_snapshot_ref=None,
                    source_version="skill-v1",
                    stage="sourced",
                    revision=2,
                )
        other = SelectionRecord(
                    store_id="store_other",
                    marketplace="amazon-ca",
                    product_name="Out of scope",
                    source_type="skill",
                    source_ref=None,
                    source_snapshot_ref=None,
                    source_version="skill-v1",
                    stage="sourced",
                    revision=1,
                )
        session.add_all([regular, dangerous, other])
        await session.flush()
        session.add_all(
            [
                SelectionQuotationVersion(
                    selection_id=regular.selection_id,
                    version=1,
                    supplier_name="Supplier missing quote",
                    unit_cost=None,
                    currency="USD",
                    value_status="missing",
                    moq=None,
                    lead_time_days=None,
                    freight_per_unit=None,
                    other_cost_per_unit=None,
                    target_price=None,
                    notes=None,
                    submission_status="submitted",
                    recorded_by=USER_IDS[RoleName.PROCUREMENT],
                ),
                SelectionQuotationVersion(
                    selection_id=regular.selection_id,
                    version=2,
                    supplier_name="Draft supplier must not export",
                    unit_cost=99,
                    currency="USD",
                    value_status="estimated",
                    moq=99,
                    lead_time_days=99,
                    freight_per_unit=99,
                    other_cost_per_unit=None,
                    target_price=None,
                    notes=None,
                    submission_status="draft",
                    recorded_by=USER_IDS[RoleName.PROCUREMENT],
                ),
                SelectionQuotationVersion(
                    selection_id=dangerous.selection_id,
                    version=1,
                    supplier_name="=Injected Supplier",
                    unit_cost=0,
                    currency="USD",
                    value_status="zero",
                    moq=12,
                    lead_time_days=0,
                    freight_per_unit=0,
                    other_cost_per_unit=None,
                    target_price=None,
                    notes=None,
                    submission_status="submitted",
                    recorded_by=USER_IDS[RoleName.PROCUREMENT],
                ),
                SelectionQuotationVersion(
                    selection_id=other.selection_id,
                    version=1,
                    supplier_name="Out of scope supplier",
                    unit_cost=10,
                    currency="CAD",
                    value_status="confirmed",
                    moq=20,
                    lead_time_days=10,
                    freight_per_unit=1,
                    other_cost_per_unit=None,
                    target_price=None,
                    notes=None,
                    submission_status="submitted",
                    recorded_by=USER_IDS[RoleName.PROCUREMENT],
                ),
            ]
        )
    try:
        yield app, audit_repository
    finally:
        await app.state.engine.dispose()


async def _login(client: httpx.AsyncClient, role: RoleName) -> dict[str, str]:
    response = await client.post(
        "/api/v1/auth/login",
        json={"username": f"{role.value}_demo", "password": PASSWORD},
    )
    assert response.status_code == 200
    cookie = response.cookies["ai_ops_session"]
    return {"Cookie": f"ai_ops_session={cookie}"}


@pytest.mark.asyncio
async def test_export_uses_frozen_scoped_snapshot_and_private_download(
    export_app: tuple[FastAPI, InMemoryAuditRepository],
) -> None:
    app, audit_repository = export_app
    async with httpx.AsyncClient(
        transport=httpx.ASGITransport(app=app), base_url="http://test"
    ) as client:
        procurement = await _login(client, RoleName.PROCUREMENT)
        operations = await _login(client, RoleName.OPERATIONS)
        creative = await _login(client, RoleName.CREATIVE)
        admin = await _login(client, RoleName.ADMIN)

        assert (
            await client.post(
                "/api/v1/selection-exports",
                headers=creative,
                json={"stage": "initial_selection"},
            )
        ).status_code == 403
        assert (
            await client.post(
                "/api/v1/selection-exports",
                headers=admin,
                json={"stage": "initial_selection"},
            )
        ).status_code == 403
        assert (
            await client.post(
                "/api/v1/selection-exports",
                headers=procurement,
                json={
                    "stage": "initial_selection",
                    "store_ids": ["store_other"],
                },
            )
        ).status_code == 403

        created = await client.post(
            "/api/v1/selection-exports",
            headers=procurement,
            json={
                "stage": "initial_selection",
                "store_ids": ["store_procurement"],
                "marketplace": "amazon-us",
            },
        )
        assert created.status_code == 201
        body = created.json()
        assert body["status"] == "completed"
        assert body["row_count"] == 2
        assert body["file_name"].endswith(".xlsx")
        assert len(body["file_checksum"]) == 64
        assert "file_path" not in body
        export_id = body["export_id"]
        snapshot_id = body["snapshot_id"]

        assert (
            await client.get(f"/api/v1/selection-exports/{export_id}", headers=operations)
        ).status_code == 404

        download = await client.get(
            f"/api/v1/selection-exports/{export_id}/download", headers=procurement
        )
        assert download.status_code == 200
        assert download.headers["content-type"].startswith(
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        )
        workbook = load_workbook(BytesIO(download.content), read_only=True, data_only=False)
        assert workbook.sheetnames == ["选品表", "导出说明"]
        sheet = workbook["选品表"]
        rows = list(sheet.iter_rows(values_only=True))
        assert rows[0] == tuple(
            column.display_name for column in InitialSelectionExporter.columns
        )
        assert len(rows) == 3
        product_index = rows[0].index("商品名称")
        source_ref_index = rows[0].index("来源引用")
        assert any(row[product_index] == "'=DANGEROUS()" for row in rows[1:])
        assert any(row[source_ref_index] is None for row in rows[1:])
        info: dict[str, object] = {
            str(row[0]): row[1]
            for row in workbook["导出说明"].iter_rows(values_only=True)
            if row[0] is not None
        }
        assert info["快照编号"] == snapshot_id
        assert info["行数"] == 2
        assert "store_procurement" in info["店铺范围"]
        assert str(info["文件校验和"]).startswith("sha256(snapshot-content):")

        async with app.state.session_factory.begin() as session:
            await session.execute(
                update(SelectionRecord)
                .where(SelectionRecord.product_name == "Regular product")
                .values(product_name="Changed after snapshot")
            )
        repeated = await client.get(
            f"/api/v1/selection-exports/{export_id}/download", headers=procurement
        )
        repeated_workbook = load_workbook(BytesIO(repeated.content), read_only=True)
        repeated_rows = list(
            repeated_workbook["选品表"].iter_rows(values_only=True)
        )
        assert "Regular product" in {row[product_index] for row in repeated_rows[1:]}
        assert "Changed after snapshot" not in {
            row[product_index] for row in repeated_rows[1:]
        }

    async with app.state.session_factory() as session:
        snapshot = await session.get(SelectionTableSnapshot, UUID(snapshot_id))
        assert snapshot is not None
        assert snapshot.row_count == 2
        assert snapshot.store_scope_json == ["store_procurement"]
        row_count = await session.scalar(
            select(func.count())
            .select_from(SelectionTableSnapshotRow)
            .where(SelectionTableSnapshotRow.snapshot_id == UUID(snapshot_id))
        )
        assert row_count == 2
    assert {
        event.action for event in audit_repository.events
    }.issuperset(
        {
            "selection.export.create",
            "selection.export.generate",
            "selection.export.download",
        }
    )


@pytest.mark.asyncio
async def test_quotation_export_freezes_latest_submitted_version_and_types(
    export_app: tuple[FastAPI, InMemoryAuditRepository],
) -> None:
    app, _ = export_app
    async with httpx.AsyncClient(
        transport=httpx.ASGITransport(app=app), base_url="http://test"
    ) as client:
        procurement = await _login(client, RoleName.PROCUREMENT)
        operations = await _login(client, RoleName.OPERATIONS)
        admin = await _login(client, RoleName.ADMIN)

        assert (
            await client.post(
                "/api/v1/selection-exports",
                headers=operations,
                json={"stage": "quotation"},
            )
        ).status_code == 403
        assert (
            await client.post(
                "/api/v1/selection-exports",
                headers=admin,
                json={"stage": "quotation"},
            )
        ).status_code == 403
        assert (
            await client.post(
                "/api/v1/selection-exports",
                headers=procurement,
                json={"stage": "quotation", "store_ids": ["store_other"]},
            )
        ).status_code == 403

        created = await client.post(
            "/api/v1/selection-exports",
            headers=procurement,
            json={"stage": "quotation", "store_ids": ["store_procurement"]},
        )
        assert created.status_code == 201
        body = created.json()
        assert body["stage"] == "quotation"
        assert body["schema_version"] == "quotation.v1"
        assert body["status"] == "completed"
        assert body["row_count"] == 2
        export_id = body["export_id"]
        snapshot_id = UUID(body["snapshot_id"])

        assert (
            await client.get(
                f"/api/v1/selection-exports/{export_id}", headers=operations
            )
        ).status_code == 404
        downloaded = await client.get(
            f"/api/v1/selection-exports/{export_id}/download", headers=procurement
        )
        assert downloaded.status_code == 200
        workbook = load_workbook(BytesIO(downloaded.content), read_only=False)
        assert workbook.sheetnames == ["采购询价表", "导出说明"]
        sheet = workbook["采购询价表"]
        rows = list(sheet.iter_rows())
        headers = [cell.value for cell in rows[0]]
        assert headers == [column.display_name for column in QuotationExporter.columns]
        values_by_product = {
            row[headers.index("商品名称")].value: row for row in rows[1:]
        }
        missing = values_by_product["Regular product"]
        zero = values_by_product["'=DANGEROUS()"]
        quote_index = headers.index("报价")
        moq_index = headers.index("最小起订量")
        freight_index = headers.index("单件运费")
        lead_time_index = headers.index("交期(天)")
        recorded_at_index = headers.index("登记时间")
        supplier_index = headers.index("供应商")
        assert missing[quote_index].value is None
        assert zero[quote_index].value == 0
        assert zero[quote_index].data_type == "n"
        assert zero[moq_index].value == 12
        assert zero[moq_index].data_type == "n"
        assert zero[freight_index].value == 0
        assert zero[freight_index].data_type == "n"
        assert zero[lead_time_index].value == 0
        assert isinstance(zero[recorded_at_index].value, datetime)
        assert zero[supplier_index].value == "'=Injected Supplier"
        assert "Draft supplier must not export" not in {
            row[supplier_index].value for row in rows[1:]
        }
        info: dict[str, object] = {
            str(row[0].value): row[1].value
            for row in workbook["导出说明"].iter_rows()
            if row[0].value is not None
        }
        assert info["询价版本范围"] == "1-1"
        assert info["草稿/正式状态"] == (
            "正式: 仅包含 submission_status=submitted"
        )
        assert info["币种说明"] == "币种: USD; 报价和单件运费按各行币种解释"

        async with app.state.session_factory.begin() as session:
            regular_id = await session.scalar(
                select(SelectionRecord.selection_id).where(
                    SelectionRecord.product_name == "Regular product"
                )
            )
            assert regular_id is not None
            session.add(
                SelectionQuotationVersion(
                    selection_id=regular_id,
                    version=3,
                    supplier_name="New supplier version",
                    unit_cost=19,
                    currency="USD",
                    value_status="confirmed",
                    moq=25,
                    lead_time_days=7,
                    freight_per_unit=2,
                    other_cost_per_unit=None,
                    target_price=None,
                    notes=None,
                    submission_status="submitted",
                    recorded_by=USER_IDS[RoleName.PROCUREMENT],
                )
            )
        old_download = await client.get(
            f"/api/v1/selection-exports/{export_id}/download", headers=procurement
        )
        old_workbook = load_workbook(BytesIO(old_download.content), read_only=True)
        old_rows = list(
            old_workbook["采购询价表"].iter_rows(values_only=True)
        )
        old_headers = old_rows[0]
        old_regular = next(
            row for row in old_rows[1:] if row[old_headers.index("商品名称")] == "Regular product"
        )
        assert old_regular[old_headers.index("报价")] is None
        assert old_regular[old_headers.index("询价版本")] == 1

    async with app.state.session_factory() as session:
        snapshot_rows = list(
            (
                await session.scalars(
                    select(SelectionTableSnapshotRow)
                    .where(SelectionTableSnapshotRow.snapshot_id == snapshot_id)
                    .order_by(SelectionTableSnapshotRow.row_position)
                )
            ).all()
        )
        assert {row.rendered_row["quotation_version"] for row in snapshot_rows} == {1}
        assert {row.rendered_row["source_version"] for row in snapshot_rows} == {
            "skill-v1"
        }


def test_exporter_preserves_zero_and_missing_and_repeats_snapshot_content(
    tmp_path: Path,
) -> None:
    exporter = InitialSelectionExporter()
    integer_column = ExportColumn("quantity", "数量", 12, "integer")
    assert exporter._typed_value(None, integer_column) is None
    assert exporter._typed_value(0, integer_column) == 0
    now = datetime(2026, 10, 9, tzinfo=UTC)
    user_id = uuid4()
    selection_id = uuid4()
    snapshot = FrozenSnapshot(
        snapshot_id=uuid4(),
        stage=SelectionExportStage.INITIAL_SELECTION,
        schema_version=exporter.schema_version,
        filters={},
        store_scope=["store_procurement"],
        requested_by=user_id,
        row_count=1,
        snapshot_checksum="a" * 64,
        created_at=now,
        rows=[
            FrozenSnapshotRow(
                row_position=1,
                selection_id=selection_id,
                selection_revision=1,
                rendered_row={
                    "selection_id": str(selection_id),
                    "store_id": "store_procurement",
                    "marketplace": "amazon-us",
                    "product_name": "Repeated row",
                    "source_type": "skill",
                    "source_ref": None,
                    "source_snapshot_ref": None,
                    "source_version": "v1",
                    "owner_id": None,
                    "stage": "sourced",
                    "revision": 0,
                    "created_at": now.isoformat(),
                    "updated_at": now.isoformat(),
                },
            )
        ],
    )
    first = exporter.export(
        snapshot,
        requested_by=user_id,
        requested_at=now,
        store_names=["采购店铺"],
        output_directory=tmp_path / "first",
    )
    second = exporter.export(
        snapshot,
        requested_by=user_id,
        requested_at=now,
        store_names=["采购店铺"],
        output_directory=tmp_path / "second",
    )
    first_rows = list(
        load_workbook(first.path, read_only=True)["选品表"].iter_rows(values_only=True)
    )
    second_rows = list(
        load_workbook(second.path, read_only=True)["选品表"].iter_rows(values_only=True)
    )
    assert first_rows == second_rows
    revision_index = first_rows[0].index("记录修订号")
    source_ref_index = first_rows[0].index("来源引用")
    assert first_rows[1][revision_index] == 0
    assert first_rows[1][source_ref_index] is None


def test_export_permission_is_not_in_any_default_role_template() -> None:
    assert all(
        Permission.SELECTION_EXPORT not in permissions
        for permissions in ROLE_PERMISSIONS.values()
    )
    assert all(
        Permission.QUOTATION_EXPORT not in permissions
        for permissions in ROLE_PERMISSIONS.values()
    )


@pytest.mark.asyncio
async def test_failed_generation_is_persisted_without_sensitive_error(
    export_app: tuple[FastAPI, InMemoryAuditRepository], tmp_path: Path
) -> None:
    original_app, audit_repository = export_app
    settings = original_app.state.settings.model_copy(
        update={"export_storage_path": tmp_path / "failed-exports"}
    )
    failing_app = create_app(
        settings,
        identity_repository=InMemoryIdentityRepository(_export_users()),
        audit_repository=audit_repository,
        admin_repository=InMemoryAdminRepository(),
        workspace_repository=InMemoryWorkspaceRepository(),
        store_repository=InMemoryStoreRepository(),
        initial_selection_exporter=FailingInitialSelectionExporter(),
    )
    try:
        async with httpx.AsyncClient(
            transport=httpx.ASGITransport(app=failing_app), base_url="http://test"
        ) as client:
            procurement = await _login(client, RoleName.PROCUREMENT)
            response = await client.post(
                "/api/v1/selection-exports",
                headers=procurement,
                json={"stage": "initial_selection"},
            )
            assert response.status_code == 201
            body = response.json()
            assert body["status"] == "failed"
            assert body["error_code"] == "EXPORT_GENERATION_FAILED"
            assert "secret" not in response.text
            persisted = await client.get(
                f"/api/v1/selection-exports/{body['export_id']}",
                headers=procurement,
            )
            assert persisted.status_code == 200
            assert persisted.json()["status"] == "failed"
            download = await client.get(
                f"/api/v1/selection-exports/{body['export_id']}/download",
                headers=procurement,
            )
            assert download.status_code == 409
    finally:
        await failing_app.state.engine.dispose()
