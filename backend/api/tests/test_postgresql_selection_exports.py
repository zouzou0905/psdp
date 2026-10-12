import os
from io import BytesIO
from pathlib import Path
from uuid import UUID, uuid4

import httpx
import pytest
from openpyxl import load_workbook
from sqlalchemy import delete, func, select, update

from app.core.config import Settings
from app.db.models import (
    GovernanceAuditLog,
    IdentityUser,
    IdentityUserPermission,
    IdentityUserStoreScope,
    SelectionExportJob,
    SelectionQuotationVersion,
    SelectionRecord,
    SelectionTableSnapshot,
    SelectionTableSnapshotRow,
)
from app.main import create_app

pytestmark = pytest.mark.skipif(
    os.environ.get("RUN_POSTGRESQL_INTEGRATION_TESTS") != "1",
    reason="Set RUN_POSTGRESQL_INTEGRATION_TESTS=1 to exercise PostgreSQL",
)


async def _login(
    client: httpx.AsyncClient, username: str, password: str
) -> dict[str, str]:
    response = await client.post(
        "/api/v1/auth/login", json={"username": username, "password": password}
    )
    response.raise_for_status()
    return {"Cookie": f"ai_ops_session={response.cookies['ai_ops_session']}"}


async def test_postgresql_export_snapshot_and_job_survive_restart() -> None:
    base_settings = Settings()
    password_secret = base_settings.demo_procurement_password
    assert password_secret is not None
    password = password_secret.get_secret_value()
    run_id = uuid4()
    storage_path = (
        Path(__file__).parents[3]
        / "tmp"
        / "exports"
        / "integration-tests"
        / str(run_id)
    )
    settings = base_settings.model_copy(update={"export_storage_path": storage_path})
    setup_app = create_app(settings)
    apps = [setup_app]
    selection_id = uuid4()
    export_id: UUID | None = None
    snapshot_id: UUID | None = None
    generated_path: Path | None = None
    permission_added = False
    async with setup_app.state.session_factory() as session:
        identity = await session.execute(
            select(IdentityUser.user_id, IdentityUserStoreScope.store_id)
            .join(
                IdentityUserStoreScope,
                IdentityUserStoreScope.user_id == IdentityUser.user_id,
            )
            .where(IdentityUser.username == "procurement_demo")
            .limit(1)
        )
        identity_row = identity.one()
        user_id: UUID = identity_row.user_id
        store_id: str = identity_row.store_id
    try:
        async with setup_app.state.session_factory.begin() as session:
            existing_permission = await session.get(
                IdentityUserPermission,
                {
                    "user_id": user_id,
                    "permission_name": "selection.export",
                },
            )
            if existing_permission is None:
                session.add(
                    IdentityUserPermission(
                        user_id=user_id, permission_name="selection.export"
                    )
                )
                permission_added = True
            await session.execute(
                update(IdentityUser)
                .where(IdentityUser.user_id == user_id)
                .values(
                    authorization_revision=IdentityUser.authorization_revision + 1
                )
            )
            session.add(
                SelectionRecord(
                    selection_id=selection_id,
                    store_id=store_id,
                    marketplace="integration-test",
                    product_name="PostgreSQL export persistence",
                    source_type="integration_test",
                    source_ref=None,
                    source_snapshot_ref="snapshot://postgresql-export",
                    source_version="test-v1",
                    stage="sourced",
                    revision=1,
                )
            )

        first_app = create_app(settings)
        apps.append(first_app)
        async with httpx.AsyncClient(
            transport=httpx.ASGITransport(app=first_app), base_url="http://test"
        ) as client:
            headers = await _login(
                client,
                "procurement_demo",
                password,
            )
            created = await client.post(
                "/api/v1/selection-exports",
                headers=headers,
                json={
                    "stage": "initial_selection",
                    "store_ids": [store_id],
                    "search": "PostgreSQL export persistence",
                },
            )
            created.raise_for_status()
            assert created.json()["status"] == "completed"
            assert created.json()["row_count"] == 1
            export_id = UUID(created.json()["export_id"])
            snapshot_id = UUID(created.json()["snapshot_id"])

        restarted_app = create_app(settings)
        apps.append(restarted_app)
        async with httpx.AsyncClient(
            transport=httpx.ASGITransport(app=restarted_app), base_url="http://test"
        ) as client:
            headers = await _login(
                client,
                "procurement_demo",
                password,
            )
            persisted = await client.get(
                f"/api/v1/selection-exports/{export_id}", headers=headers
            )
            persisted.raise_for_status()
            assert persisted.json()["snapshot_id"] == str(snapshot_id)
            downloaded = await client.get(
                f"/api/v1/selection-exports/{export_id}/download", headers=headers
            )
            downloaded.raise_for_status()
            workbook = load_workbook(BytesIO(downloaded.content), read_only=True)
            assert workbook.sheetnames == ["选品表", "导出说明"]
            assert downloaded.content.startswith(b"PK")

        async with setup_app.state.session_factory() as session:
            assert export_id is not None
            job = await session.get(SelectionExportJob, export_id)
            assert job is not None
            generated_path = Path(job.file_path) if job.file_path else None
            assert job.file_checksum
            assert snapshot_id is not None
            snapshot = await session.get(SelectionTableSnapshot, snapshot_id)
            assert snapshot is not None
            assert snapshot.snapshot_checksum
            frozen_count = await session.scalar(
                select(func.count())
                .select_from(SelectionTableSnapshotRow)
                .where(SelectionTableSnapshotRow.snapshot_id == snapshot_id)
            )
            assert frozen_count == 1
    finally:
        async with setup_app.state.session_factory.begin() as session:
            if export_id is not None:
                await session.execute(
                    delete(GovernanceAuditLog).where(
                        GovernanceAuditLog.resource_id == str(export_id)
                    )
                )
                await session.execute(
                    delete(SelectionExportJob).where(
                        SelectionExportJob.export_id == export_id
                    )
                )
            if snapshot_id is not None:
                await session.execute(
                    delete(SelectionTableSnapshotRow).where(
                        SelectionTableSnapshotRow.snapshot_id == snapshot_id
                    )
                )
                await session.execute(
                    delete(SelectionTableSnapshot).where(
                        SelectionTableSnapshot.snapshot_id == snapshot_id
                    )
                )
            await session.execute(
                delete(SelectionRecord).where(
                    SelectionRecord.selection_id == selection_id
                )
            )
            if permission_added:
                await session.execute(
                    delete(IdentityUserPermission).where(
                        IdentityUserPermission.user_id == user_id,
                        IdentityUserPermission.permission_name == "selection.export",
                    )
                )
            await session.execute(
                update(IdentityUser)
                .where(IdentityUser.user_id == user_id)
                .values(
                    authorization_revision=IdentityUser.authorization_revision + 1
                )
            )
        if generated_path is not None:
            generated_path.unlink(missing_ok=True)
        for app in apps:
            await app.state.engine.dispose()


async def test_postgresql_quotation_export_freezes_version() -> None:
    settings = Settings()
    password_secret = settings.demo_procurement_password
    assert password_secret is not None
    password = password_secret.get_secret_value()
    app = create_app(settings)
    selection_id = uuid4()
    export_id: UUID | None = None
    snapshot_id: UUID | None = None
    generated_path: Path | None = None
    permission_added = False
    async with app.state.session_factory() as session:
        identity = await session.execute(
            select(IdentityUser.user_id, IdentityUserStoreScope.store_id)
            .join(
                IdentityUserStoreScope,
                IdentityUserStoreScope.user_id == IdentityUser.user_id,
            )
            .where(IdentityUser.username == "procurement_demo")
            .limit(1)
        )
        identity_row = identity.one()
        user_id: UUID = identity_row.user_id
        store_id: str = identity_row.store_id
    try:
        async with app.state.session_factory.begin() as session:
            permission = await session.get(
                IdentityUserPermission,
                {"user_id": user_id, "permission_name": "quotation.export"},
            )
            if permission is None:
                session.add(
                    IdentityUserPermission(
                        user_id=user_id, permission_name="quotation.export"
                    )
                )
                permission_added = True
            await session.execute(
                update(IdentityUser)
                .where(IdentityUser.user_id == user_id)
                .values(
                    authorization_revision=IdentityUser.authorization_revision + 1
                )
            )
            session.add(
                SelectionRecord(
                    selection_id=selection_id,
                    store_id=store_id,
                    marketplace="integration-test",
                    product_name="PostgreSQL quotation export",
                    source_type="integration_test",
                    source_ref=None,
                    source_snapshot_ref="snapshot://quotation-export",
                    source_version="source-v1",
                    stage="quoted",
                    revision=1,
                )
            )
            await session.flush()
            session.add(
                SelectionQuotationVersion(
                    selection_id=selection_id,
                    version=1,
                    supplier_name="PostgreSQL supplier",
                    unit_cost=0,
                    currency="USD",
                    value_status="zero",
                    moq=10,
                    lead_time_days=5,
                    freight_per_unit=None,
                    other_cost_per_unit=None,
                    target_price=None,
                    notes=None,
                    submission_status="submitted",
                    recorded_by=user_id,
                )
            )
        async with httpx.AsyncClient(
            transport=httpx.ASGITransport(app=app), base_url="http://test"
        ) as client:
            headers = await _login(client, "procurement_demo", password)
            response = await client.post(
                "/api/v1/selection-exports",
                headers=headers,
                json={
                    "stage": "quotation",
                    "store_ids": [store_id],
                    "search": "PostgreSQL quotation export",
                },
            )
            response.raise_for_status()
            body = response.json()
            assert body["status"] == "completed"
            assert body["schema_version"] == "quotation.v1"
            export_id = UUID(body["export_id"])
            snapshot_id = UUID(body["snapshot_id"])
            downloaded = await client.get(
                f"/api/v1/selection-exports/{export_id}/download", headers=headers
            )
            downloaded.raise_for_status()
            workbook = load_workbook(BytesIO(downloaded.content), read_only=True)
            assert workbook.sheetnames == ["采购询价表", "导出说明"]
            rows = list(workbook["采购询价表"].iter_rows(values_only=True))
            assert rows[1][rows[0].index("报价")] == 0
            assert rows[1][rows[0].index("询价版本")] == 1
        async with app.state.session_factory() as session:
            assert export_id is not None
            job = await session.get(SelectionExportJob, export_id)
            assert job is not None
            generated_path = Path(job.file_path) if job.file_path else None
            assert snapshot_id is not None
            snapshot_row = await session.scalar(
                select(SelectionTableSnapshotRow).where(
                    SelectionTableSnapshotRow.snapshot_id == snapshot_id
                )
            )
            assert snapshot_row is not None
            assert snapshot_row.rendered_row["source_version"] == "source-v1"
            assert snapshot_row.rendered_row["quotation_version"] == 1
    finally:
        async with app.state.session_factory.begin() as session:
            if export_id is not None:
                await session.execute(
                    delete(GovernanceAuditLog).where(
                        GovernanceAuditLog.resource_id == str(export_id)
                    )
                )
                await session.execute(
                    delete(SelectionExportJob).where(
                        SelectionExportJob.export_id == export_id
                    )
                )
            if snapshot_id is not None:
                await session.execute(
                    delete(SelectionTableSnapshotRow).where(
                        SelectionTableSnapshotRow.snapshot_id == snapshot_id
                    )
                )
                await session.execute(
                    delete(SelectionTableSnapshot).where(
                        SelectionTableSnapshot.snapshot_id == snapshot_id
                    )
                )
            await session.execute(
                delete(SelectionRecord).where(
                    SelectionRecord.selection_id == selection_id
                )
            )
            if permission_added:
                await session.execute(
                    delete(IdentityUserPermission).where(
                        IdentityUserPermission.user_id == user_id,
                        IdentityUserPermission.permission_name == "quotation.export",
                    )
                )
            await session.execute(
                update(IdentityUser)
                .where(IdentityUser.user_id == user_id)
                .values(
                    authorization_revision=IdentityUser.authorization_revision + 1
                )
            )
        if generated_path is not None:
            generated_path.unlink(missing_ok=True)
        await app.state.engine.dispose()
